from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlmodel import Session

from population_data.models import (
    AttentionOrQualityQuestion,
    BinaryQuestion,
    CategoricalQuestion,
    DataKind,
    DemographicQuestion,
    DerivedQuestion,
    ExperimentalConditionQuestion,
    LikertQuestion,
    MultiSelectQuestion,
    NumericQuestion,
    Option,
    OrdinalQuestion,
    Person,
    PopulationDataset,
    Question,
    QuestionRole,
    Response,
    ResponseValue,
    Scale,
    SemanticDifferentialQuestion,
    SourceRef,
    Study,
    TextQuestion,
    ValueKind,
)
from services import population_serialization_service
from services.population_serialization_service import PopulationDatasetImportResult


REQUIRED_SHEETS = {"study", "questions", "responses"}
OPTIONAL_SHEETS = {"options"}

STUDY_HEADERS = ["field", "value"]
QUESTION_HEADERS = [
    "code",
    "parent_code",
    "source_type",
    "data_kind",
    "role",
    "question_text",
    "item_text",
    "wording",
    "scale_min",
    "scale_max",
    "scale_min_label",
    "scale_max_label",
    "scale_step",
    "required",
    "randomized",
    "max_selections",
    "sum_to",
    "short_name",
    "constructs",
    "missing_codes",
    "reverse_scored",
    "harmonized_variable_id",
    "notes",
]
OPTION_HEADERS = ["question_code", "value", "label", "order", "is_missing", "notes"]
RESPONSE_PERSON_ID_COLUMN = "person_id"


@dataclass(frozen=True)
class ValidationIssue:
    location: str
    message: str
    severity: str = "error"


def read_population_xlsx(path: Path, *, source_file_path: str | None = None) -> PopulationDataset:
    workbook = load_workbook(path, read_only=True, data_only=True)
    issues = _workbook_shape_issues(workbook.sheetnames)
    if issues:
        raise ValueError(_issue_summary(issues))

    source_path = source_file_path or str(path)
    study = _read_study(workbook["study"], path, source_path)
    option_rows = _read_options(workbook["options"]) if "options" in workbook.sheetnames else {}
    questions = _read_questions(workbook["questions"], study.id, source_path, option_rows)
    question_ids = {question.id for question in questions}
    unknown_option_codes = sorted(set(option_rows) - question_ids)
    if unknown_option_codes:
        raise ValueError(f"Options reference unknown question codes: {', '.join(unknown_option_codes)}")
    people = _read_responses(workbook["responses"], study.id, source_path, question_ids, option_rows)
    dataset = PopulationDataset(
        studies=[study],
        questions=questions,
        people=people,
        metadata={
            "source_format": "strict_population_xlsx",
            "source_file": str(path),
        },
    )
    issues = validate_population_dataset(dataset)
    if issues:
        raise ValueError(_issue_summary(issues))
    return dataset


def is_strict_population_xlsx(path: Path) -> bool:
    workbook = load_workbook(path, read_only=True, data_only=True)
    return REQUIRED_SHEETS.issubset(set(workbook.sheetnames))


def write_population_xlsx(dataset: PopulationDataset, path: Path) -> None:
    issues = validate_population_dataset(dataset)
    if issues:
        raise ValueError(_issue_summary(issues))
    if len(dataset.studies) != 1:
        raise ValueError("XLSX export requires exactly one study.")

    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_study(workbook.create_sheet("study"), dataset.studies[0])
    _write_questions(workbook.create_sheet("questions"), dataset.questions)
    _write_options(workbook.create_sheet("options"), dataset.questions)
    _write_responses(workbook.create_sheet("responses"), dataset)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def import_population_xlsx(
    session: Session,
    path: Path,
    *,
    skip_existing: bool = True,
    source_file_path: str | None = None,
) -> PopulationDatasetImportResult:
    dataset = read_population_xlsx(path, source_file_path=source_file_path)
    return population_serialization_service.import_population_dataset(session, dataset, skip_existing=skip_existing)


def export_population_xlsx(
    session: Session,
    *,
    study_id: int,
    path: Path,
    include_people: bool = True,
) -> PopulationDataset:
    dataset = population_serialization_service.export_population_dataset(
        session,
        study_ids=[study_id],
        include_people=include_people,
    )
    write_population_xlsx(dataset, path)
    return dataset


def validate_population_dataset(dataset: PopulationDataset) -> list[ValidationIssue]:
    issues: list[ValidationIssue] = []
    if len(dataset.studies) != 1:
        issues.append(ValidationIssue("studies", "XLSX datasets must contain exactly one study."))
        return issues

    study_id = dataset.studies[0].id
    question_ids: set[str] = set()
    for question in dataset.questions:
        if question.id in question_ids:
            issues.append(ValidationIssue(f"questions.{question.id}", "Duplicate question code."))
        question_ids.add(question.id)
        if question.study_id != study_id:
            issues.append(ValidationIssue(f"questions.{question.id}", "Question study_id does not match workbook study."))

    person_ids: set[str] = set()
    option_values_by_question = _option_values_by_question(dataset.questions)
    for person in dataset.people:
        if person.id in person_ids:
            issues.append(ValidationIssue(f"responses.{person.id}", "Duplicate person_id."))
        person_ids.add(person.id)
        if person.study_id != study_id:
            issues.append(ValidationIssue(f"responses.{person.id}", "Person study_id does not match workbook study."))
        seen_response_ids: set[str] = set()
        for response in person.responses:
            if response.question_id not in question_ids:
                issues.append(
                    ValidationIssue(
                        f"responses.{person.id}.{response.question_id}",
                        "Response references a question code not defined in questions.",
                    )
                )
            if response.question_id in seen_response_ids:
                issues.append(
                    ValidationIssue(
                        f"responses.{person.id}.{response.question_id}",
                        "Duplicate response for person/question.",
                    )
                )
            seen_response_ids.add(response.question_id)
            allowed_values = option_values_by_question.get(response.question_id)
            if allowed_values is not None and not response.value.is_missing and response.value.normalized not in allowed_values:
                issues.append(
                    ValidationIssue(
                        f"responses.{person.id}.{response.question_id}",
                        f"Response value {response.value.normalized!r} is not listed in options.",
                    )
                )
    return issues


def _workbook_shape_issues(sheet_names: Iterable[str]) -> list[ValidationIssue]:
    names = set(sheet_names)
    issues = [
        ValidationIssue("workbook", f"Missing required sheet: {sheet_name}")
        for sheet_name in sorted(REQUIRED_SHEETS - names)
    ]
    unknown = names - REQUIRED_SHEETS - OPTIONAL_SHEETS
    for sheet_name in sorted(unknown):
        issues.append(ValidationIssue(f"sheet.{sheet_name}", "Unknown sheet will not be imported.", "warning"))
    return [issue for issue in issues if issue.severity == "error"]


def _read_study(sheet: Worksheet, path: Path, source_path: str) -> Study:
    values = {
        str(row[0]).strip(): _cell_value(row[1])
        for row in _data_rows(sheet)
        if row and row[0] not in (None, "")
    }
    study_id = _required(values, "study_external_id", "study")
    return Study(
        id=str(study_id),
        name=str(values.get("name") or study_id),
        source_file=str(values.get("source_file") or Path(source_path).name),
        source_sheet=_optional_str(values.get("source_sheet")),
        description=_optional_str(values.get("description")),
        collected_at=_optional_str(values.get("collected_at")),
        metadata={
            key: value
            for key, value in values.items()
            if key not in {"study_external_id", "name", "source_file", "source_sheet", "description", "collected_at"}
        },
    )


def _read_questions(
    sheet: Worksheet,
    study_id: str,
    source_path: str,
    option_rows: dict[str, list[Option]],
) -> list[Question]:
    rows = _dict_rows(sheet)
    questions: list[Question] = []
    seen: set[str] = set()
    for index, row in enumerate(rows, start=2):
        code = _required(row, "code", f"questions row {index}")
        if code in seen:
            raise ValueError(f"Duplicate question code in questions sheet: {code}")
        seen.add(str(code))
        kind = _enum(DataKind, row.get("data_kind"), DataKind.UNKNOWN, f"questions row {index} data_kind")
        role = _enum(QuestionRole, row.get("role"), QuestionRole.UNKNOWN, f"questions row {index} role")
        question = _question_for_kind(
            kind,
            {
                "id": str(code),
                "study_id": study_id,
                "wording": str(row.get("wording") or row.get("question_text") or code),
                "source": SourceRef(
                    file_path=source_path,
                    sheet_name=sheet.title,
                    row_index=index,
                    column_index=1,
                    column_name="code",
                ),
                "short_name": _optional_str(row.get("short_name")) or str(code),
                "role": role,
                "constructs": _json_list(row.get("constructs")),
                "scale": _scale_from_question_row(row, option_rows.get(str(code), [])),
                "missing_codes": _json_list(row.get("missing_codes")),
                "reverse_scored": bool(row.get("reverse_scored") or False),
                "harmonized_variable_id": _optional_str(row.get("harmonized_variable_id")),
                "metadata": _question_metadata(row),
            },
        )
        questions.append(question)
    return questions


def _read_options(sheet: Worksheet) -> dict[str, list[Option]]:
    options: dict[str, list[Option]] = {}
    for index, row in enumerate(_dict_rows(sheet), start=2):
        question_code = _required(row, "question_code", f"options row {index}")
        options.setdefault(str(question_code), []).append(
            Option(
                value=_cell_value(row.get("value")),
                label=_optional_str(row.get("label")),
                order=_optional_int(row.get("order")),
                is_missing=bool(row.get("is_missing") or False),
            )
        )
    return options


def _read_responses(
    sheet: Worksheet,
    study_id: str,
    source_path: str,
    question_ids: set[str],
    option_rows: dict[str, list[Option]],
) -> list[Person]:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    if not headers or headers[0] != RESPONSE_PERSON_ID_COLUMN:
        raise ValueError("responses sheet must start with a person_id column.")
    response_question_ids = [header for header in headers[1:] if header]
    unknown = [header for header in response_question_ids if header not in question_ids]
    if unknown:
        raise ValueError(f"Response columns missing from questions sheet: {', '.join(unknown)}")

    people: list[Person] = []
    seen_people: set[str] = set()
    for row_index, row in enumerate(rows[1:], start=2):
        if not any(value is not None for value in row):
            continue
        person_id = str(row[0]).strip() if row and row[0] is not None else ""
        if not person_id:
            raise ValueError(f"responses row {row_index} is missing person_id.")
        if person_id in seen_people:
            raise ValueError(f"Duplicate person_id in responses sheet: {person_id}")
        seen_people.add(person_id)
        responses: list[Response] = []
        for column_index, question_id in enumerate(response_question_ids, start=2):
            raw = row[column_index - 1] if column_index - 1 < len(row) else None
            responses.append(
                Response(
                    question_id=question_id,
                    value=_response_value(raw, option_rows.get(question_id, [])),
                    source=SourceRef(
                        file_path=source_path,
                        sheet_name=sheet.title,
                        row_index=row_index,
                        column_index=column_index,
                        column_name=question_id,
                    ),
                )
            )
        people.append(
            Person(
                id=person_id,
                study_id=study_id,
                source_person_id=person_id,
                source=SourceRef(file_path=source_path, sheet_name=sheet.title, row_index=row_index),
                responses=responses,
            )
        )
    return people


def _write_study(sheet: Worksheet, study: Study) -> None:
    sheet.append(STUDY_HEADERS)
    rows = {
        "study_external_id": study.id,
        "name": study.name,
        "description": study.description,
        "source_file": study.source_file,
        "source_sheet": study.source_sheet,
        "collected_at": study.collected_at,
        **study.metadata,
    }
    for key, value in rows.items():
        sheet.append([key, _xlsx_value(value)])


def _write_questions(sheet: Worksheet, questions: list[Question]) -> None:
    sheet.append(QUESTION_HEADERS)
    for question in questions:
        metadata = question.metadata or {}
        scale = question.scale
        sheet.append(
            [
                question.id,
                metadata.get("parent_code") or question.id,
                metadata.get("source_type") or "",
                question.kind.value,
                question.role.value,
                metadata.get("question_text") or question.wording,
                metadata.get("item_text") or "",
                question.wording,
                scale.min_value if scale else None,
                scale.max_value if scale else None,
                scale.min_label if scale else None,
                scale.max_label if scale else None,
                scale.step if scale else None,
                _constraint(metadata, "required"),
                _constraint(metadata, "randomized"),
                _constraint(metadata, "max_selections"),
                _constraint(metadata, "sum_to"),
                question.short_name,
                json.dumps(question.constructs, ensure_ascii=False),
                json.dumps(question.missing_codes, ensure_ascii=False),
                question.reverse_scored,
                question.harmonized_variable_id,
                metadata.get("notes") or "",
            ]
        )


def _write_options(sheet: Worksheet, questions: list[Question]) -> None:
    sheet.append(OPTION_HEADERS)
    for question in questions:
        if question.scale is None:
            continue
        for option in question.scale.options:
            sheet.append([question.id, _xlsx_value(option.value), option.label, option.order, option.is_missing, ""])


def _write_responses(sheet: Worksheet, dataset: PopulationDataset) -> None:
    question_ids = [question.id for question in dataset.questions]
    sheet.append([RESPONSE_PERSON_ID_COLUMN, *question_ids])
    for person in dataset.people:
        values_by_question = {response.question_id: response.value.raw for response in person.responses}
        sheet.append([person.id, *[_xlsx_value(values_by_question.get(question_id)) for question_id in question_ids]])


def _question_for_kind(kind: DataKind, base: dict[str, Any]) -> Question:
    if kind == DataKind.NUMERIC:
        return NumericQuestion(**base)
    if kind == DataKind.ORDINAL:
        return OrdinalQuestion(**base)
    if kind == DataKind.LIKERT:
        return LikertQuestion(**base)
    if kind == DataKind.SEMANTIC_DIFFERENTIAL:
        return SemanticDifferentialQuestion(**base)
    if kind == DataKind.BINARY:
        return BinaryQuestion(**base)
    if kind == DataKind.EXPERIMENTAL_CONDITION:
        return ExperimentalConditionQuestion(**base)
    if kind == DataKind.MULTI_SELECT:
        return MultiSelectQuestion(**base)
    if kind == DataKind.TEXT:
        return TextQuestion(**base)
    if kind == DataKind.DERIVED:
        return DerivedQuestion(**base)
    if kind == DataKind.ATTENTION_OR_QUALITY:
        return AttentionOrQualityQuestion(**base)
    if kind == DataKind.DEMOGRAPHIC:
        return DemographicQuestion(**base)
    return CategoricalQuestion(**base)


def _scale_from_question_row(row: dict[str, Any], options: list[Option]) -> Scale | None:
    has_scale = any(
        row.get(key) not in (None, "")
        for key in ("scale_min", "scale_max", "scale_min_label", "scale_max_label", "scale_step")
    )
    if not has_scale and not options:
        return None
    return Scale(
        min_value=_optional_float(row.get("scale_min")),
        max_value=_optional_float(row.get("scale_max")),
        min_label=_optional_str(row.get("scale_min_label")),
        max_label=_optional_str(row.get("scale_max_label")),
        step=_optional_float(row.get("scale_step")),
        options=options,
    )


def _response_value(value: Any, options: list[Option]) -> ResponseValue:
    if value is None or value == "":
        return ResponseValue(kind=ValueKind.MISSING, raw=value, normalized=None, is_missing=True)
    normalized = _cell_value(value)
    label = None
    is_missing = False
    for option in options:
        if option.value == normalized:
            label = option.label
            is_missing = option.is_missing
            break
    if isinstance(normalized, bool):
        kind = ValueKind.BOOLEAN
    elif isinstance(normalized, int) and not isinstance(normalized, bool):
        kind = ValueKind.INTEGER
    elif isinstance(normalized, float):
        kind = ValueKind.NUMBER
    elif isinstance(normalized, datetime):
        kind = ValueKind.DATETIME
        normalized = normalized.isoformat()
    elif isinstance(normalized, date):
        kind = ValueKind.DATE
        normalized = normalized.isoformat()
    else:
        kind = ValueKind.STRING
        normalized = str(normalized)
    return ResponseValue(kind=kind, raw=_cell_value(value), normalized=normalized, label=label, is_missing=is_missing)


def _option_values_by_question(questions: list[Question]) -> dict[str, set[Any] | None]:
    values: dict[str, set[Any] | None] = {}
    for question in questions:
        if question.scale and question.scale.options:
            values[question.id] = {option.value for option in question.scale.options}
        else:
            values[question.id] = None
    return values


def _dict_rows(sheet: Worksheet) -> list[dict[str, Any]]:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    dicts: list[dict[str, Any]] = []
    for row in rows[1:]:
        if not any(value is not None for value in row):
            continue
        dicts.append({headers[index]: _cell_value(value) for index, value in enumerate(row) if index < len(headers) and headers[index]})
    return dicts


def _data_rows(sheet: Worksheet) -> list[tuple[Any, ...]]:
    return [row for row in sheet.iter_rows(values_only=True) if any(value is not None for value in row)]


def _question_metadata(row: dict[str, Any]) -> dict[str, Any]:
    constraints = {
        "required": row.get("required"),
        "randomized": row.get("randomized"),
        "max_selections": row.get("max_selections"),
        "sum_to": row.get("sum_to"),
    }
    return {
        "parent_code": row.get("parent_code") or row.get("code"),
        "source_type": row.get("source_type"),
        "question_text": row.get("question_text"),
        "item_text": row.get("item_text"),
        "constraints": {key: value for key, value in constraints.items() if value not in (None, "")},
        "notes": row.get("notes") or "",
    }


def _constraint(metadata: dict[str, Any], key: str) -> Any:
    constraints = metadata.get("constraints") or {}
    return constraints.get(key)


def _json_list(value: Any) -> list[Any]:
    if value in (None, ""):
        return []
    if isinstance(value, list):
        return value
    if isinstance(value, str):
        parsed = json.loads(value)
        if not isinstance(parsed, list):
            raise ValueError(f"Expected JSON list, got: {value}")
        return parsed
    return [value]


def _cell_value(value: Any) -> Any:
    if isinstance(value, datetime | date):
        return value.isoformat()
    return value


def _xlsx_value(value: Any) -> Any:
    if isinstance(value, dict | list):
        return json.dumps(value, ensure_ascii=False)
    return value


def _optional_str(value: Any) -> str | None:
    return None if value in (None, "") else str(value)


def _optional_int(value: Any) -> int | None:
    if value in (None, ""):
        return None
    return int(value)


def _optional_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    return float(value)


def _required(mapping: dict[str, Any], key: str, location: str) -> Any:
    value = mapping.get(key)
    if value in (None, ""):
        raise ValueError(f"Missing required field {key!r} in {location}.")
    return value


def _enum(enum_type, value: Any, default, location: str):
    if value in (None, ""):
        return default
    try:
        return enum_type(value)
    except ValueError as exc:
        raise ValueError(f"Invalid {location}: {value}") from exc


def _issue_summary(issues: list[ValidationIssue]) -> str:
    return "; ".join(f"{issue.location}: {issue.message}" for issue in issues)
