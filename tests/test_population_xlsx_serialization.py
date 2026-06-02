from __future__ import annotations

from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook
from sqlmodel import Session, SQLModel, create_engine, select
from sqlalchemy.pool import StaticPool

from models import PopulationPersonRecord, PopulationQuestionRecord, PopulationResponseRecord
from repositories import population_repo
from services.population_xlsx_serialization_service import (
    export_population_xlsx,
    import_population_xlsx,
    read_population_xlsx,
    write_population_xlsx,
)


def strict_workbook(
    path: Path,
    *,
    question_rows: list[list[object]] | None = None,
    option_rows: list[list[object]] | None = None,
    response_rows: list[list[object]] | None = None,
    include_options: bool = True,
) -> None:
    workbook = Workbook()
    study = workbook.active
    study.title = "study"
    study.append(["field", "value"])
    study.append(["study_external_id", "mini_food_study"])
    study.append(["name", "Mini Food Study"])
    study.append(["source_file", "mini.xlsx"])
    study.append(["source_sheet", "responses"])

    questions = workbook.create_sheet("questions")
    questions.append(
        [
            "code",
            "parent_code",
            "source_type",
            "data_kind",
            "role",
            "question_text",
            "item_text",
            "wording",
            "scale_min",
            "scale_max",
            "scale_min_label",
            "scale_max_label",
            "scale_step",
            "required",
            "randomized",
            "max_selections",
            "sum_to",
            "notes",
        ]
    )
    for row in question_rows or [
        [
            "q1_1",
            "q1",
            "Matrix",
            "likert",
            "belief",
            "How much do you agree?",
            "I buy sustainable food.",
            "How much do you agree? I buy sustainable food.",
            1,
            7,
            "strongly disagree",
            "strongly agree",
            1,
            True,
            False,
            None,
            None,
            "",
        ],
        [
            "gender",
            "gender",
            "Single",
            "demographic",
            "demographic",
            "What is your gender?",
            "",
            "What is your gender?",
            None,
            None,
            "",
            "",
            None,
            True,
            False,
            None,
            None,
            "",
        ],
    ]:
        questions.append(row)

    if include_options:
        options = workbook.create_sheet("options")
        options.append(["question_code", "value", "label", "order", "is_missing", "notes"])
        for row in option_rows or [
            ["q1_1", 1, "strongly disagree", 1, False, ""],
            ["q1_1", 7, "strongly agree", 7, False, ""],
            ["gender", 1, "Woman", 1, False, ""],
            ["gender", 2, "Man", 2, False, ""],
            ["gender", 99, "Prefer not to say", 3, True, ""],
        ]:
            options.append(row)

    responses = workbook.create_sheet("responses")
    for row in response_rows or [
        ["person_id", "q1_1", "gender"],
        ["P001", 7, 1],
        ["P002", 1, 99],
    ]:
        responses.append(row)

    workbook.save(path)


class PopulationXlsxSerializationTests(unittest.TestCase):
    def test_strict_workbook_reads_complete_dataset(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "strict.xlsx"
            strict_workbook(path)

            dataset = read_population_xlsx(path)

        self.assertEqual(len(dataset.studies), 1)
        self.assertEqual(dataset.studies[0].id, "mini_food_study")
        self.assertEqual(len(dataset.questions), 2)
        self.assertEqual(dataset.questions[0].id, "q1_1")
        self.assertEqual(dataset.questions[0].scale.max_label, "strongly agree")
        self.assertEqual(dataset.questions[1].scale.options[2].is_missing, True)
        self.assertEqual(len(dataset.people), 2)
        self.assertEqual(dataset.people[0].responses[0].question_id, "q1_1")
        self.assertEqual(dataset.people[0].responses[0].value.normalized, 7)

    def test_written_workbook_round_trips(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.xlsx"
            exported = Path(tmp) / "exported.xlsx"
            strict_workbook(source)
            dataset = read_population_xlsx(source)

            write_population_xlsx(dataset, exported)
            round_tripped = read_population_xlsx(exported)

        self.assertEqual(round_tripped.studies[0].id, dataset.studies[0].id)
        self.assertEqual([question.id for question in round_tripped.questions], [question.id for question in dataset.questions])
        self.assertEqual([person.id for person in round_tripped.people], [person.id for person in dataset.people])

    def test_missing_required_sheet_fails(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "missing.xlsx"
            workbook = Workbook()
            workbook.active.title = "study"
            workbook.save(path)

            with self.assertRaisesRegex(ValueError, "Missing required sheet"):
                read_population_xlsx(path)

    def test_response_column_missing_from_questions_fails(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "unknown_response.xlsx"
            strict_workbook(path, response_rows=[["person_id", "q1_1", "unknown"], ["P001", 7, 1]])

            with self.assertRaisesRegex(ValueError, "Response columns missing from questions"):
                read_population_xlsx(path)

    def test_option_for_unknown_question_fails(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "unknown_option.xlsx"
            strict_workbook(path, option_rows=[["unknown", 1, "Unknown", 1, False, ""]])

            with self.assertRaisesRegex(ValueError, "Options reference unknown question codes"):
                read_population_xlsx(path)

    def test_duplicate_question_code_fails(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "duplicate_question.xlsx"
            duplicate_row = [
                "q1_1",
                "q1",
                "Matrix",
                "likert",
                "belief",
                "How much do you agree?",
                "Again.",
                "How much do you agree? Again.",
                1,
                7,
                "strongly disagree",
                "strongly agree",
                1,
                True,
                False,
                None,
                None,
                "",
            ]
            strict_workbook(path, question_rows=[duplicate_row, duplicate_row])

            with self.assertRaisesRegex(ValueError, "Duplicate question code"):
                read_population_xlsx(path)

    def test_duplicate_person_id_fails(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "duplicate_person.xlsx"
            strict_workbook(path, response_rows=[["person_id", "q1_1", "gender"], ["P001", 7, 1], ["P001", 1, 2]])

            with self.assertRaisesRegex(ValueError, "Duplicate person_id"):
                read_population_xlsx(path)

    def test_invalid_data_kind_fails(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "invalid_kind.xlsx"
            row = [
                "q1_1",
                "q1",
                "Matrix",
                "not_a_kind",
                "belief",
                "How much do you agree?",
                "I buy sustainable food.",
                "How much do you agree? I buy sustainable food.",
                1,
                7,
                "strongly disagree",
                "strongly agree",
                1,
                True,
                False,
                None,
                None,
                "",
            ]
            strict_workbook(path, question_rows=[row])

            with self.assertRaisesRegex(ValueError, "Invalid questions row 2 data_kind"):
                read_population_xlsx(path)

    def test_import_population_xlsx_persists_dataset_and_skips_existing(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "strict.xlsx"
            strict_workbook(path)
            engine = create_engine(
                "sqlite://",
                connect_args={"check_same_thread": False},
                poolclass=StaticPool,
            )
            SQLModel.metadata.create_all(engine)

            with Session(engine) as session:
                first = import_population_xlsx(session, path)
                session.commit()
                study = population_repo.get_study_by_external_id(session, "mini_food_study")

                second = import_population_xlsx(session, path)
                session.commit()

                self.assertIsNotNone(study)
                self.assertEqual(first.studies_created, 1)
                self.assertEqual(first.questions_created, 2)
                self.assertEqual(first.people_created, 2)
                self.assertEqual(first.responses_created, 4)
                self.assertEqual(second.studies_created, 0)
                self.assertEqual(second.questions_created, 0)
                self.assertEqual(second.people_created, 0)
                self.assertEqual(second.responses_created, 0)
                self.assertEqual(second.people_skipped, 2)
                self.assertEqual(population_repo.count_questions(session, int(study.id)), 2)
                self.assertEqual(population_repo.count_people(session, int(study.id)), 2)
                self.assertEqual(len(list(session.exec(select(PopulationQuestionRecord)))), 2)
                self.assertEqual(len(list(session.exec(select(PopulationPersonRecord)))), 2)
                self.assertEqual(len(list(session.exec(select(PopulationResponseRecord)))), 4)

    def test_export_population_xlsx_writes_db_study(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "strict.xlsx"
            exported = Path(tmp) / "db_export.xlsx"
            strict_workbook(source)
            engine = create_engine(
                "sqlite://",
                connect_args={"check_same_thread": False},
                poolclass=StaticPool,
            )
            SQLModel.metadata.create_all(engine)

            with Session(engine) as session:
                import_population_xlsx(session, source)
                session.commit()
                study = population_repo.get_study_by_external_id(session, "mini_food_study")

                export_population_xlsx(session, study_id=int(study.id), path=exported)

            round_tripped = read_population_xlsx(exported)
            self.assertEqual(round_tripped.studies[0].id, "mini_food_study")
            self.assertEqual(len(round_tripped.questions), 2)
            self.assertEqual(len(round_tripped.people), 2)


if __name__ == "__main__":
    unittest.main()
