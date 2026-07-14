from __future__ import annotations

import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

import openpyxl
from pydantic import TypeAdapter

from models import (
    AttentionOrQualityQuestion,
    BinaryQuestion,
    CategoricalQuestion,
    DataKind,
    DemographicQuestion,
    DerivedQuestion,
    ExperimentalConditionQuestion,
    LikertQuestion,
    NumericQuestion,
    OrdinalQuestion,
    Person,
    PopulationDataset,
    Question,
    QuestionRole,
    Response,
    ResponseValue,
    Scale,
    SemanticDifferentialQuestion,
    SourceRef,
    Study,
    TextQuestion,
    ValueKind,
)


ROOT = Path(__file__).resolve().parent
OUTPUT_PATH = ROOT / "outputs" / "population_data.json"

QUESTION_CATALOG_MARKER = "question-items"
QUESTION_ADAPTER = TypeAdapter(Question)


def slugify(value: str, max_length: int = 80) -> str:
    value = value.strip().lower()
    value = re.sub(r"[^a-z0-9]+", "_", value)
    value = re.sub(r"_+", "_", value).strip("_")
    return (value[:max_length].strip("_") or "item")


def compact_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value)).strip()


def iter_xlsx_files(root: Path) -> Iterable[Path]:
    ignored_parts = {".venv", "outputs", "__pycache__"}
    for path in sorted(root.rglob("*.xlsx")):
        if ignored_parts.intersection(path.parts):
            continue
        yield path


def make_study_id(path: Path, sheet_name: str) -> str:
    return f"{slugify(path.relative_to(ROOT).as_posix())}__{slugify(sheet_name)}"


def make_question_id(study_id: str, column_index: int, wording: str) -> str:
    return f"{study_id}__q{column_index:03d}_{slugify(wording, 48)}"


def make_catalog_question_id(path: Path, sheet_name: str, row_index: int, wording: str) -> str:
    file_slug = slugify(path.relative_to(ROOT).as_posix(), 50)
    return f"{file_slug}__{slugify(sheet_name, 30)}__row{row_index:03d}_{slugify(wording, 45)}"


def jsonable(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def response_value(value: Any) -> ResponseValue:
    if value is None:
        return ResponseValue(kind=ValueKind.MISSING, raw=None, normalized=None, is_missing=True)
    if isinstance(value, bool):
        return ResponseValue(kind=ValueKind.BOOLEAN, raw=value, normalized=value)
    if isinstance(value, int) and not isinstance(value, bool):
        return ResponseValue(kind=ValueKind.INTEGER, raw=value, normalized=value)
    if isinstance(value, float):
        return ResponseValue(kind=ValueKind.NUMBER, raw=value, normalized=value)
    if isinstance(value, datetime):
        return ResponseValue(kind=ValueKind.DATETIME, raw=value.isoformat(), normalized=value.isoformat())
    if isinstance(value, date):
        return ResponseValue(kind=ValueKind.DATE, raw=value.isoformat(), normalized=value.isoformat())
    text = compact_text(value)
    if text == "":
        return ResponseValue(kind=ValueKind.MISSING, raw="", normalized=None, is_missing=True)
    return ResponseValue(kind=ValueKind.STRING, raw=text, normalized=text)


def infer_role(wording: str) -> QuestionRole:
    text = wording.lower()
    demographic_terms = [
        "age",
        "alder",
        "gender",
        "køn",
        "region",
        "postcode",
        "postal code",
        "kommune",
        "municipality",
        "education",
        "uddannelse",
        "household",
        "hushold",
        "income",
        "indkomst",
    ]
    quality_terms = ["duration", "status", "completion", "exclude", "attention", "threeminutes"]
    condition_terms = ["condition", "gruppe", "group", "route", "produkt", "product", "viser hvilket"]
    behavior_terms = ["how often", "hvor ofte", "spiser", "eat", "thrown away", "smidt ud", "købt", "buy"]

    if any(term in text for term in quality_terms):
        return QuestionRole.QUALITY_CONTROL
    if any(term in text for term in demographic_terms):
        return QuestionRole.DEMOGRAPHIC
    if any(term in text for term in condition_terms):
        return QuestionRole.EXPERIMENTAL_CONDITION
    if "average" in text or "_avg" in text or "score" in text:
        return QuestionRole.DERIVED
    if any(term in text for term in behavior_terms):
        return QuestionRole.BEHAVIOR
    if "likely" in text or "sandsynligt" in text or "synes du om" in text:
        return QuestionRole.PREFERENCE
    if "agree" in text or "enig" in text or "opfattelse" in text:
        return QuestionRole.BELIEF
    return QuestionRole.UNKNOWN


def infer_constructs(wording: str) -> list[str]:
    text = wording.lower()
    constructs: list[str] = []
    rules = {
        "environmental_concern": ["environment", "miljø", "klima", "bæredygtig", "nature", "natur"],
        "meat_reduction": ["kød", "meat", "vegetar", "plant", "plante"],
        "health_orientation": ["health", "sund", "helbred", "protein"],
        "food_waste": ["waste", "thrown away", "smidt ud", "leftover", "bread"],
        "trust": ["trust", "troværdig", "honest", "kompetent"],
        "food_neophilia": ["try new", "nye fødevarer", "usædvanlige"],
        "product_attitude": ["bad/good", "dårligt", "godt", "liking", "synes du om"],
        "purchase_intention": ["buy", "købe", "sandsynligt"],
    }
    for construct, terms in rules.items():
        if any(term in text for term in terms):
            constructs.append(construct)
    return constructs


def infer_question_kind(wording: str, values: list[Any]) -> DataKind:
    role = infer_role(wording)
    text = wording.lower()
    non_missing = [value for value in values if value is not None and compact_text(value) != ""]
    unique_values = set(json.dumps(jsonable(value), ensure_ascii=False) for value in non_missing)

    if role == QuestionRole.DEMOGRAPHIC:
        return DataKind.DEMOGRAPHIC
    if role == QuestionRole.QUALITY_CONTROL:
        return DataKind.ATTENTION_OR_QUALITY
    if role == QuestionRole.EXPERIMENTAL_CONDITION:
        return DataKind.EXPERIMENTAL_CONDITION
    if role == QuestionRole.DERIVED:
        return DataKind.DERIVED
    if "bad/good" in text or "dårligt" in text or "useless/useful" in text or "dumt" in text:
        return DataKind.SEMANTIC_DIFFERENTIAL
    if "agree" in text or "enig" in text:
        return DataKind.LIKERT
    if non_missing and all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in non_missing):
        if len(unique_values) <= 2:
            return DataKind.BINARY
        if 3 <= len(unique_values) <= 12:
            return DataKind.ORDINAL
        return DataKind.NUMERIC
    if len(unique_values) <= 20:
        return DataKind.CATEGORICAL
    return DataKind.TEXT


def make_scale(values: list[Any], raw_description: str | None = None) -> Scale | None:
    numbers = [value for value in values if isinstance(value, (int, float)) and not isinstance(value, bool)]
    if numbers:
        return Scale(
            min_value=float(min(numbers)),
            max_value=float(max(numbers)),
            raw_description=raw_description,
        )
    if raw_description:
        return Scale(raw_description=raw_description)
    return None


def build_question(
    *,
    id: str,
    study_id: str,
    kind: DataKind,
    wording: str,
    source: SourceRef,
    values: list[Any] | None = None,
    scale_description: str | None = None,
) -> Question:
    values = values or []
    role = infer_role(wording)
    base = {
        "id": id,
        "study_id": study_id,
        "wording": wording,
        "source": source,
        "short_name": source.column_name,
        "role": role,
        "constructs": infer_constructs(wording),
        "scale": make_scale(values, scale_description),
        "metadata": {},
    }

    if kind == DataKind.NUMERIC:
        return NumericQuestion(**base)
    if kind == DataKind.ORDINAL:
        return OrdinalQuestion(**base)
    if kind == DataKind.LIKERT:
        return LikertQuestion(**base)
    if kind == DataKind.SEMANTIC_DIFFERENTIAL:
        return SemanticDifferentialQuestion(**base)
    if kind == DataKind.BINARY:
        return BinaryQuestion(**base)
    if kind == DataKind.EXPERIMENTAL_CONDITION:
        return ExperimentalConditionQuestion(**base)
    if kind == DataKind.TEXT:
        return TextQuestion(**base)
    if kind == DataKind.DERIVED:
        return DerivedQuestion(**base)
    if kind == DataKind.ATTENTION_OR_QUALITY:
        return AttentionOrQualityQuestion(**base)
    if kind == DataKind.DEMOGRAPHIC:
        return DemographicQuestion(**base)
    return CategoricalQuestion(**base)


def read_sheet_rows(ws: Any) -> list[tuple[Any, ...]]:
    return [row for row in ws.iter_rows(values_only=True) if any(cell is not None for cell in row)]


def ingest_data_workbook(path: Path) -> tuple[list[Study], list[Question], list[Person]]:
    studies: list[Study] = []
    questions: list[Question] = []
    people: list[Person] = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    for ws in wb.worksheets:
        rows = read_sheet_rows(ws)
        if len(rows) < 2:
            continue
        headers = [compact_text(value) if value is not None else "" for value in rows[0]]
        if not any(headers):
            continue

        study_id = make_study_id(path, ws.title)
        studies.append(
            Study(
                id=study_id,
                name=ws.title,
                source_file=path.relative_to(ROOT).as_posix(),
                source_sheet=ws.title,
                metadata={"row_count": len(rows) - 1, "column_count": len(headers)},
            )
        )

        question_ids_by_col: dict[int, str] = {}
        for col_idx, header in enumerate(headers, start=1):
            if not header:
                continue
            values = [row[col_idx - 1] if col_idx - 1 < len(row) else None for row in rows[1:]]
            kind = infer_question_kind(header, values)
            question_id = make_question_id(study_id, col_idx, header)
            question_ids_by_col[col_idx] = question_id
            questions.append(
                build_question(
                    id=question_id,
                    study_id=study_id,
                    kind=kind,
                    wording=header,
                    source=SourceRef(
                        file_path=path.relative_to(ROOT).as_posix(),
                        sheet_name=ws.title,
                        row_index=1,
                        column_index=col_idx,
                        column_name=header,
                    ),
                    values=values,
                )
            )

        for row_offset, row in enumerate(rows[1:], start=2):
            source_person_id = row[0] if row else None
            person_id = f"{study_id}__person_{row_offset - 1:05d}"
            responses: list[Response] = []
            attributes: dict[str, ResponseValue] = {}
            for col_idx, question_id in question_ids_by_col.items():
                raw_value = row[col_idx - 1] if col_idx - 1 < len(row) else None
                value = response_value(raw_value)
                response = Response(
                    question_id=question_id,
                    value=value,
                    source=SourceRef(
                        file_path=path.relative_to(ROOT).as_posix(),
                        sheet_name=ws.title,
                        row_index=row_offset,
                        column_index=col_idx,
                        column_name=headers[col_idx - 1],
                    ),
                )
                responses.append(response)
                header = headers[col_idx - 1]
                role = infer_role(header)
                if role in {QuestionRole.DEMOGRAPHIC, QuestionRole.METADATA, QuestionRole.QUALITY_CONTROL}:
                    attributes[slugify(header)] = value

            people.append(
                Person(
                    id=person_id,
                    study_id=study_id,
                    source_person_id=jsonable(source_person_id),
                    attributes=attributes,
                    responses=responses,
                    source=SourceRef(
                        file_path=path.relative_to(ROOT).as_posix(),
                        sheet_name=ws.title,
                        row_index=row_offset,
                    ),
                )
            )

    return studies, questions, people


def ingest_question_catalog(path: Path) -> tuple[list[Study], list[Question], list[Person]]:
    studies: list[Study] = []
    questions: list[Question] = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    for ws in wb.worksheets:
        rows = read_sheet_rows(ws)
        if len(rows) < 3:
            continue

        study_id = make_study_id(path, ws.title)
        studies.append(
            Study(
                id=study_id,
                name=f"Question catalog: {ws.title}",
                source_file=path.relative_to(ROOT).as_posix(),
                source_sheet=ws.title,
                description=compact_text(rows[0][0]) if rows and rows[0] and rows[0][0] else None,
                metadata={"catalog_rows": len(rows)},
            )
        )

        for row_idx, row in enumerate(rows[2:], start=3):
            wording = compact_text(row[0]) if len(row) > 0 and row[0] is not None else ""
            if not wording:
                continue
            scale_description = compact_text(row[1]) if len(row) > 1 and row[1] is not None else None
            if scale_description is None and len(row) <= 2:
                continue
            if scale_description is None and not re.search(r"[?]", wording):
                continue
            kind = infer_question_kind(wording, [])
            if scale_description and re.search(r"\b0-1\b|\b1-2\b|yes/no|ja/nej", scale_description.lower()):
                kind = DataKind.BINARY
            elif scale_description and re.search(r"\b1-\d+\b", scale_description.lower()):
                kind = DataKind.LIKERT if ("agree" in wording.lower() or "enig" in wording.lower()) else DataKind.ORDINAL
            elif scale_description and "open" in scale_description.lower():
                kind = DataKind.TEXT

            questions.append(
                build_question(
                    id=make_catalog_question_id(path, ws.title, row_idx, wording),
                    study_id=study_id,
                    kind=kind,
                    wording=wording,
                    source=SourceRef(
                        file_path=path.relative_to(ROOT).as_posix(),
                        sheet_name=ws.title,
                        row_index=row_idx,
                        column_index=1,
                        column_name="Question-item",
                    ),
                    values=[],
                    scale_description=scale_description,
                )
            )

    return studies, questions, []


def build_dataset() -> PopulationDataset:
    studies: list[Study] = []
    questions: list[Question] = []
    people: list[Person] = []
    source_files: list[str] = []

    for path in iter_xlsx_files(ROOT):
        source_files.append(path.relative_to(ROOT).as_posix())
        if QUESTION_CATALOG_MARKER in path.name.lower():
            file_studies, file_questions, file_people = ingest_question_catalog(path)
        else:
            file_studies, file_questions, file_people = ingest_data_workbook(path)
        studies.extend(file_studies)
        questions.extend(file_questions)
        people.extend(file_people)

    return PopulationDataset(
        studies=studies,
        questions=questions,
        people=people,
        metadata={
            "source_files": source_files,
            "xlsx_file_count": len(source_files),
            "study_count": len(studies),
            "question_count": len(questions),
            "person_count": len(people),
            "encoding": "utf-8",
            "note": "Raw source values are preserved as UTF-8 text. Harmonization is intentionally deferred.",
        },
    )


def main() -> None:
    dataset = build_dataset()
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(
        dataset.model_dump_json(exclude_none=True, ensure_ascii=False),
        encoding="utf-8",
    )
    print(f"Wrote {OUTPUT_PATH.relative_to(ROOT)}")
    print(
        f"{len(dataset.studies)} studies, "
        f"{len(dataset.questions)} questions, "
        f"{len(dataset.people)} people"
    )


if __name__ == "__main__":
    main()
